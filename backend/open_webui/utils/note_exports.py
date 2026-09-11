"""Small, deterministic exports for private Markdown notes."""

from __future__ import annotations

import re
from io import BytesIO
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


_HEADING = re.compile(r'^(#{1,6})\s+(.+?)\s*$')
_BULLET = re.compile(r'^\s*[-*+]\s+(.+?)\s*$')
_ORDERED = re.compile(r'^\s*\d+[.)]\s+(.+?)\s*$')
_IMAGE = re.compile(r'!\[([^\]]*)\]\((/api/v1/files/([^/?\s)]+)/content(?:\?[^)]*)?)\)')
_MARKUP = re.compile(r'(`+|\*\*|__|\*|_|~~)')
_FONT_FAMILY = 'Microsoft YaHei'


def internal_note_image_ids(markdown: str) -> list[tuple[str, str]]:
    """Return only internal file references; exports never fetch remote image URLs."""
    return [(match.group(1) or 'Image', match.group(3)) for match in _IMAGE.finditer(markdown or '')]


def _plain_text(value: str) -> str:
    value = _MARKUP.sub('', value)
    return re.sub(r'\[([^\]]+)\]\([^)]*\)', r'\1', value).strip()


def _set_document_font(style) -> None:
    style.font.name = _FONT_FAMILY
    style._element.rPr.rFonts.set(qn('w:ascii'), _FONT_FAMILY)
    style._element.rPr.rFonts.set(qn('w:hAnsi'), _FONT_FAMILY)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), _FONT_FAMILY)


def _clear_title_bottom_border(paragraph) -> None:
    paragraph_properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'nil')
    borders.append(bottom)
    paragraph_properties.append(borders)


def _add_markdown_body(document: Document, markdown: str) -> None:
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        if paragraph_lines:
            document.add_paragraph(_plain_text(' '.join(paragraph_lines)))
            paragraph_lines.clear()

    for raw_line in (markdown or '').splitlines():
        line = raw_line.strip()
        heading = _HEADING.match(line)
        if heading:
            flush_paragraph()
            document.add_heading(_plain_text(heading.group(2)), level=min(len(heading.group(1)), 3))
            continue
        bullet = _BULLET.match(line)
        if bullet:
            flush_paragraph()
            document.add_paragraph(_plain_text(bullet.group(1)), style='List Bullet')
            continue
        ordered = _ORDERED.match(line)
        if ordered:
            flush_paragraph()
            document.add_paragraph(_plain_text(ordered.group(1)), style='List Number')
            continue
        if not line:
            flush_paragraph()
            continue
        if line.startswith('|') and line.endswith('|'):
            if set(line.replace('|', '').strip()) <= {'-', ':', ' '}:
                continue
            flush_paragraph()
            document.add_paragraph('  |  '.join(_plain_text(cell.strip()) for cell in line.strip('|').split('|')))
            continue
        if line.startswith('!['):
            continue
        paragraph_lines.append(line)

    flush_paragraph()


def render_note_docx(title: str, markdown: str, images: Iterable[tuple[str, bytes]] = ()) -> bytes:
    """Render a portable DOCX without trusting remote content or document metadata."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    normal = document.styles['Normal']
    _set_document_font(normal)
    normal.font.size = Pt(10.5)
    for style_name in ('Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        style = document.styles[style_name]
        _set_document_font(style)
        style.font.color.rgb = RGBColor(0, 0, 0)

    title_paragraph = document.add_paragraph(style='Title')
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_paragraph.add_run((title or 'Untitled note').strip())
    _clear_title_bottom_border(title_paragraph)
    _add_markdown_body(document, markdown)

    for alt_text, image_bytes in images:
        if not image_bytes:
            continue
        try:
            document.add_picture(BytesIO(image_bytes), width=Inches(5.8))
            caption = document.add_paragraph(alt_text or 'Image')
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:
            # A malformed source image must not make the note or its text unavailable.
            continue

    output = BytesIO()
    document.save(output)
    return output.getvalue()
