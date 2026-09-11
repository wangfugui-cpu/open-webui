"""Security and content checks for private YanChuan note downloads."""

from __future__ import annotations

import base64
import json
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from fastapi import HTTPException
from open_webui.models.notes import NoteModel
from open_webui.models.users import UserModel
from open_webui.routers import notes as notes_router
from open_webui.tools import builtin


def _user(user_id: str) -> UserModel:
    return UserModel(
        id=user_id,
        email=f'{user_id}@example.test',
        name=user_id,
        role='user',
        last_active_at=0,
        updated_at=0,
        created_at=0,
    )


def _note() -> NoteModel:
    return NoteModel(
        id='private-note',
        user_id='owner',
        title='冲突整理说明',
        data={
            'content': {
                'md': '# 结论\n\n资料 A 与资料 B 的上线日期不同。\n\n- 确认负责人\n- 确认日期\n'
            }
        },
        access_grants=[],
        created_at=0,
        updated_at=0,
    )


async def _body(response) -> bytes:
    return b''.join([chunk async for chunk in response.body_iterator])


@pytest.mark.asyncio
async def test_private_note_downloads_are_access_checked_and_docx_contains_note(monkeypatch):
    note = _note()

    async def get_note(*_args, **_kwargs):
        return note

    async def permitted(*_args, **_kwargs):
        return True

    async def no_grant(*_args, **_kwargs):
        return False

    monkeypatch.setattr(notes_router.Notes, 'get_note_by_id', get_note)
    monkeypatch.setattr(notes_router, 'has_permission', permitted)
    monkeypatch.setattr(notes_router.AccessGrants, 'has_access', no_grant)

    markdown_response = await notes_router.download_note_by_id('private-note', 'md', user=_user('owner'), db=None)
    assert markdown_response.headers['content-type'].startswith('text/markdown')
    assert 'attachment;' in markdown_response.headers['content-disposition']
    assert '资料 A' in (await _body(markdown_response)).decode('utf-8')

    docx_response = await notes_router.download_note_by_id('private-note', 'docx', user=_user('owner'), db=None)
    assert docx_response.headers['content-type'].startswith(
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    document = Document(BytesIO(await _body(docx_response)))
    text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
    assert '冲突整理说明' in text
    assert '资料 A 与资料 B 的上线日期不同。' in text
    assert '确认负责人' in text

    with pytest.raises(HTTPException) as denied:
        await notes_router.download_note_by_id('private-note', 'md', user=_user('other-user'), db=None)
    assert denied.value.status_code == 403


@pytest.mark.asyncio
async def test_docx_export_keeps_only_authorized_internal_images(monkeypatch):
    note = _note()
    image = base64.b64decode(
        'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9JcYQAAAAASUVORK5CYII='
    )

    async def get_note(*_args, **_kwargs):
        return note

    async def permitted(*_args, **_kwargs):
        return True

    async def authorized_images(*_args, **_kwargs):
        return [('编辑后的蓝杯', image)]

    monkeypatch.setattr(notes_router.Notes, 'get_note_by_id', get_note)
    monkeypatch.setattr(notes_router, 'has_permission', permitted)
    monkeypatch.setattr(notes_router, '_note_export_images', authorized_images)

    response = await notes_router.download_note_by_id('private-note', 'docx', user=_user('owner'), db=None)
    exported = await _body(response)
    with ZipFile(BytesIO(exported)) as archive:
        assert any(name.startswith('word/media/') for name in archive.namelist())
    document = Document(BytesIO(exported))
    assert '编辑后的蓝杯' in '\n'.join(paragraph.text for paragraph in document.paragraphs)


def test_docx_export_turns_markdown_schedule_into_a_real_table():
    from open_webui.utils.note_exports import render_note_docx

    rendered = render_note_docx(
        '安排说明',
        '| 时间 | 事项 |\n| --- | --- |\n| 周一 | 收集资料 |\n| 周二 | 确认冲突 |',
    )
    document = Document(BytesIO(rendered))

    assert len(document.tables) == 1
    assert [cell.text for cell in document.tables[0].rows[0].cells] == ['时间', '事项']
    assert [cell.text for cell in document.tables[0].rows[1].cells] == ['周一', '收集资料']


@pytest.mark.asyncio
async def test_export_note_tool_returns_private_product_download_link(monkeypatch):
    note = _note()

    async def get_note(*_args, **_kwargs):
        return note

    async def read_access(_note, user):
        return user.get('id') == 'owner'

    monkeypatch.setattr(builtin.Notes, 'get_note_by_id', get_note)
    monkeypatch.setattr(builtin, '_has_read_access_to_note', read_access)

    exported = await builtin.export_note(
        'private-note',
        'docx',
        __request__=object(),
        __user__={'id': 'owner', 'role': 'user'},
    )
    assert json.loads(exported) == {
        'status': 'success',
        'id': 'private-note',
        'title': '冲突整理说明',
        'format': 'docx',
        'download_url': '/notes/private-note?download=docx',
    }

    denied = await builtin.export_note(
        'private-note',
        'docx',
        __request__=object(),
        __user__={'id': 'other-user', 'role': 'user'},
    )
    assert json.loads(denied)['code'] == 'access_denied'
